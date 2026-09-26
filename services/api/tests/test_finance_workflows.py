"""Regression coverage for the finance bounded-context decomposition (Sprint 110).

Each test walks a critical finance workflow end to end through the public API so
the extracted domains (contributions, receipts, custody, expenses, budgeting,
reminders, reporting) keep their observable behaviour.
"""

import uuid
from io import BytesIO

import pytest
from helpers import create_tenant_with_user, create_user_for_tenant, login
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.asyncio


async def _finance_people(db_session: AsyncSession, suffix: str) -> dict:
    admin = await create_tenant_with_user(db_session, suffix)
    treasurer = await create_user_for_tenant(
        db_session,
        tenant_id=admin["tenant"].id,
        email=f"treasurer.{suffix}@test.org",
        password="TreasurerPass123!",
        display_name="Treasurer Finance",
        role_code="treasurer",
        profile_type="staff",
    )
    member = await create_user_for_tenant(
        db_session,
        tenant_id=admin["tenant"].id,
        email=f"member.{suffix}@test.org",
        password="MemberPass123!",
        display_name="Member Finance",
        role_code="member",
        profile_type="member",
        member_code=f"FIN-{suffix[:6].upper()}",
    )
    await db_session.commit()
    return {"admin": admin, "treasurer": treasurer, "member": member}


async def test_contribution_records_payment_and_statement_lifecycle(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    suffix = uuid.uuid4().hex[:6]
    people = await _finance_people(db_session, suffix)
    admin, treasurer, member = people["admin"], people["treasurer"], people["member"]
    treasurer_token = await login(client, treasurer["user"].email, treasurer["password"], admin["tenant"].slug)
    headers = {"Authorization": f"Bearer {treasurer_token}"}

    created = await client.post(
        "/api/v1/contributions/",
        json={
            "membership_profile_id": str(member["profile"].id),
            "year": 2026,
            "expected_amount": "120.00",
            "paid_amount": "0.00",
            "currency": "EUR",
            "status": "pending",
        },
        headers=headers,
    )
    assert created.status_code == 201, created.text
    record = created.json()
    assert record["balance"] == "120.00"

    listed = await client.get("/api/v1/contributions/", params={"year": 2026}, headers=headers)
    assert listed.status_code == 200
    assert [row["id"] for row in listed.json()] == [record["id"]]

    by_member = await client.get(
        f"/api/v1/contributions/by-member/{member['profile'].id}", headers=headers
    )
    assert by_member.status_code == 200
    assert by_member.json()[0]["id"] == record["id"]

    payment = await client.post(
        "/api/v1/contributions/payments",
        json={
            "contribution_record_id": record["id"],
            "amount": "50.00",
            "payment_method": "bank_transfer",
            "reference": "PAY-1",
        },
        headers=headers,
    )
    assert payment.status_code == 201, payment.text

    refreshed = await client.get(f"/api/v1/contributions/{record['id']}", headers=headers)
    assert refreshed.json()["paid_amount"] == "50.00"
    assert refreshed.json()["balance"] == "70.00"

    tenant_payments = await client.get("/api/v1/contributions/payments", headers=headers)
    assert len(tenant_payments.json()) == 1
    per_contribution = await client.get(
        f"/api/v1/contributions/{record['id']}/payments", headers=headers
    )
    assert per_contribution.json()[0]["id"] == payment.json()["id"]

    summary = await client.get(
        "/api/v1/contributions/summary", params={"year": 2026}, headers=headers
    )
    assert summary.json() == {
        "total_count": 1,
        "total_expected": "120.00",
        "total_paid": "50.00",
        "total_balance": "70.00",
    }

    statement = await client.get(
        f"/api/v1/memberships/{member['profile'].id}/statement", headers=headers
    )
    assert statement.status_code == 200, statement.text
    assert statement.json()["summary"]["total_paid"] == "50.00"
    assert statement.json()["contributions"][0]["id"] == record["id"]

    updated = await client.patch(
        f"/api/v1/contributions/{record['id']}",
        json={"expected_amount": "150.00"},
        headers=headers,
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["expected_amount"] == "150.00"

    admin_token = await login(client, admin["user"].email, admin["password"], admin["tenant"].slug)
    deleted = await client.delete(
        f"/api/v1/contributions/{record['id']}",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert deleted.status_code == 204
    remaining = await client.get("/api/v1/contributions/", headers=headers)
    assert remaining.json() == []


async def test_receipt_rejection_requires_note_and_stays_out_of_budget(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    suffix = uuid.uuid4().hex[:6]
    people = await _finance_people(db_session, suffix)
    admin, treasurer = people["admin"], people["treasurer"]
    treasurer_token = await login(client, treasurer["user"].email, treasurer["password"], admin["tenant"].slug)
    treasurer_headers = {"Authorization": f"Bearer {treasurer_token}"}

    declared = await client.post(
        "/api/v1/contributions/receipt-declarations",
        json={"income_type": "donation", "source_name": "Anonymous donor", "amount": "30.00"},
        headers=treasurer_headers,
    )
    assert declared.status_code == 201, declared.text
    declaration_id = declared.json()["id"]
    submitted = await client.post(
        f"/api/v1/contributions/receipt-declarations/{declaration_id}/submit",
        headers=treasurer_headers,
    )
    assert submitted.status_code == 200, submitted.text

    missing_note = await client.post(
        f"/api/v1/contributions/receipt-declarations/{declaration_id}/process",
        json={"action": "rejected"},
        headers=treasurer_headers,
    )
    assert missing_note.status_code == 422

    rejected = await client.post(
        f"/api/v1/contributions/receipt-declarations/{declaration_id}/process",
        json={"action": "rejected", "note": "No supporting receipt"},
        headers=treasurer_headers,
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["status"] == "rejected"
    assert rejected.json()["processing_note"] == "No supporting receipt"
    assert rejected.json()["cash_handover_status"] is None

    budget = await client.get(
        "/api/v1/contributions/annual-budget", params={"year": 2026}, headers=treasurer_headers
    )
    assert budget.status_code == 200, budget.text
    assert budget.json()["income_total"] == "0.00"


async def test_expense_budget_and_finance_exports_end_to_end(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    suffix = uuid.uuid4().hex[:6]
    people = await _finance_people(db_session, suffix)
    admin, treasurer, member = people["admin"], people["treasurer"], people["member"]
    treasurer_token = await login(client, treasurer["user"].email, treasurer["password"], admin["tenant"].slug)
    headers = {"Authorization": f"Bearer {treasurer_token}"}

    created = await client.post(
        "/api/v1/contributions/",
        json={
            "membership_profile_id": str(member["profile"].id),
            "year": 2026,
            "expected_amount": "60.00",
            "paid_amount": "0.00",
            "currency": "EUR",
            "status": "pending",
        },
        headers=headers,
    )
    assert created.status_code == 201, created.text
    record = created.json()

    payment = await client.post(
        "/api/v1/contributions/payments",
        json={"contribution_record_id": record["id"], "amount": "60.00", "payment_method": "cash"},
        headers=headers,
    )
    assert payment.status_code == 201, payment.text

    expense = await client.post(
        "/api/v1/contributions/expenses",
        json={
            "category": "administration",
            "amount": "10.00",
            "currency": "EUR",
            "spent_at": "2026-06-01T12:00:00Z",
            "description": "Office supplies",
            "payee": "Paper shop",
            "payment_method": "bank_transfer",
        },
        headers=headers,
    )
    assert expense.status_code == 201, expense.text

    budget = await client.get(
        "/api/v1/contributions/annual-budget", params={"year": 2026}, headers=headers
    )
    assert budget.status_code == 200, budget.text
    body = budget.json()
    assert body["income_total"] == "60.00"
    assert body["expense_total"] == "10.00"
    assert body["available_balance"] == "50.00"
    assert body["income_by_category"] == [
        {"category": "membership_contribution", "amount": "60.00"},
        {"category": "donation", "amount": "0.00"},
        {"category": "sponsorship", "amount": "0.00"},
        {"category": "tournament_proceeds", "amount": "0.00"},
        {"category": "disciplinary_payment", "amount": "0.00"},
        {"category": "other_income", "amount": "0.00"},
    ]
    assert body["expenses_by_category"][0] == {"category": "sport_equipment", "amount": "0.00"}
    assert body["expenses_by_category"][4] == {"category": "administration", "amount": "10.00"}
    assert len(body["recent_expenses"]) == 1

    admin_token = await login(client, admin["user"].email, admin["password"], admin["tenant"].slug)
    contributions_csv = await client.get(
        "/api/v1/contributions/export",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert contributions_csv.status_code == 200
    assert str(member["profile"].id) in contributions_csv.text
    assert "60.00" in contributions_csv.text

    report_csv = await client.get(
        "/api/v1/contributions/report/export",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert report_csv.status_code == 200
    assert "payment_count" in report_csv.text

    xlsx = await client.get(
        "/api/v1/contributions/report/export/xlsx",
        params={"year": 2026},
        headers=headers,
    )
    assert xlsx.status_code == 200, xlsx.text
    workbook = load_workbook(BytesIO(xlsx.content))
    sheet = workbook.active
    assert sheet.title == "Cotisations"
    assert "Rapport des cotisations 2026" in sheet["A1"].value
    assert sheet.max_row >= 5
    assert sheet.cell(row=5, column=1).value == member["profile"].member_code

    pdf = await client.get(
        "/api/v1/contributions/report/export/pdf",
        params={"year": 2026},
        headers=headers,
    )
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")

    whatsapp = await client.get(
        "/api/v1/contributions/report/export/whatsapp",
        params={"year": 2026},
        headers=headers,
    )
    assert whatsapp.status_code == 200, whatsapp.text
    assert "Member Finance" in whatsapp.text
    assert "Total dû" in whatsapp.text

    statement = await client.get(
        f"/api/v1/memberships/{member['profile'].id}/statement", headers=headers
    )
    assert statement.status_code == 200, statement.text
    assert statement.json()["summary"]["total_paid"] == "60.00"
