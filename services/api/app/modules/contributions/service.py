import io
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from uuid import UUID

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.import_export import ImportResult, ImportRowError, generate_csv, parse_csv
from app.modules.audit.service import AuditService
from app.modules.contributions.models import (
    ContributionReceiptStatus,
    ContributionStatus,
    ReminderDeliveryStatus,
)
from app.modules.contributions.repository import ContributionRepository
from app.modules.contributions.schemas import (
    ContributionReceiptDeclarationCreate,
    ContributionReceiptDeclarationProcess,
    ContributionReceiptDeclarationResponse,
    ContributionReceiptDeclarationUpdate,
    ContributionRecordCreate,
    ContributionRecordResponse,
    ContributionRecordUpdate,
    ContributionReminderBatchRequest,
    ContributionReminderBatchResponse,
    ContributionReminderResponse,
    ContributionReminderSendRequest,
    PaymentRecordCreate,
    PaymentRecordResponse,
)
from app.modules.disciplinary.models import DisciplinaryRecord
from app.modules.membership.models import MembershipProfile
from app.modules.membership.repository import MembershipRepository
from app.modules.tenancy.models import Tenant
from app.providers.notifications.base import NotificationDispatchResult, NotificationProvider


class ContributionService:
    def __init__(self, db: AsyncSession) -> None:
        self._db = db
        self._repo = ContributionRepository(db)
        self._audit = AuditService(db)
        self._notification_providers: list[NotificationProvider] = []

    def with_notification_providers(
        self, providers: list[NotificationProvider]
    ) -> "ContributionService":
        self._notification_providers = providers
        return self

    async def create_contribution(
        self,
        tenant_id: UUID,
        data: ContributionRecordCreate,
        *,
        actor_user_id: UUID | None = None,
    ) -> ContributionRecordResponse:
        record = await self._repo.create_contribution(
            tenant_id, data.model_dump()
        )
        await self._audit.record_event(
            tenant_id=tenant_id,
            actor_user_id=actor_user_id,
            action="create",
            entity_type="contribution_record",
            entity_id=record.id,
            module_key="contributions",
            details={
                "membership_profile_id": record.membership_profile_id,
                "year": record.year,
                "expected_amount": str(record.expected_amount),
                "paid_amount": str(record.paid_amount),
            },
        )
        await self._db.commit()
        return ContributionRecordResponse.model_validate(record)

    async def get_contribution(
        self, tenant_id: UUID, contribution_id: UUID
    ) -> ContributionRecordResponse:
        record = await self._repo.get_contribution(tenant_id, contribution_id)
        if not record:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Contribution record not found",
            )
        return ContributionRecordResponse.model_validate(record)

    async def list_contributions(
        self, tenant_id: UUID, year: int | None = None
    ) -> list[ContributionRecordResponse]:
        records = await self._repo.list_by_tenant(tenant_id, year)
        return [ContributionRecordResponse.model_validate(r) for r in records]

    async def list_member_contributions(
        self, tenant_id: UUID, profile_id: UUID
    ) -> list[ContributionRecordResponse]:
        records = await self._repo.list_by_profile(tenant_id, profile_id)
        return [ContributionRecordResponse.model_validate(r) for r in records]

    async def update_contribution(
        self,
        tenant_id: UUID,
        contribution_id: UUID,
        data: ContributionRecordUpdate,
        *,
        actor_user_id: UUID | None = None,
    ) -> ContributionRecordResponse:
        record = await self._repo.update_contribution(
            tenant_id, contribution_id, data.model_dump(exclude_unset=True)
        )
        if not record:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Contribution record not found",
            )
        await self._audit.record_event(
            tenant_id=tenant_id,
            actor_user_id=actor_user_id,
            action="update",
            entity_type="contribution_record",
            entity_id=record.id,
            module_key="contributions",
            details={"changes": data.model_dump(exclude_unset=True)},
        )
        await self._db.commit()
        return ContributionRecordResponse.model_validate(record)

    async def delete_contribution(
        self,
        tenant_id: UUID,
        contribution_id: UUID,
        *,
        actor_user_id: UUID | None = None,
    ) -> None:
        existing = await self._repo.get_contribution(tenant_id, contribution_id)
        deleted = await self._repo.delete_contribution(tenant_id, contribution_id)
        if not deleted:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Contribution record not found",
            )
        await self._audit.record_event(
            tenant_id=tenant_id,
            actor_user_id=actor_user_id,
            action="delete",
            entity_type="contribution_record",
            entity_id=contribution_id,
            module_key="contributions",
            details={
                "year": existing.year if existing else None,
                "membership_profile_id": existing.membership_profile_id if existing else None,
            },
        )
        await self._db.commit()

    async def record_payment(
        self,
        tenant_id: UUID,
        data: PaymentRecordCreate,
        *,
        actor_user_id: UUID | None = None,
    ) -> PaymentRecordResponse:
        contrib = await self._repo.get_contribution(
            tenant_id, data.contribution_record_id
        )
        if not contrib:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Contribution record not found",
            )
        payload = data.model_dump()
        if payload.get("paid_at") is None:
            from datetime import datetime
            payload["paid_at"] = datetime.now(UTC)
        payment = await self._repo.create_payment(tenant_id, payload)
        new_paid = contrib.paid_amount + data.amount
        await self._repo.update_contribution(
            tenant_id, data.contribution_record_id, {"paid_amount": new_paid}
        )
        await self._audit.record_event(
            tenant_id=tenant_id,
            actor_user_id=actor_user_id,
            action="payment_recorded",
            entity_type="payment_record",
            entity_id=payment.id,
            module_key="contributions",
            details={
                "contribution_record_id": contrib.id,
                "amount": str(payment.amount),
                "method": payment.payment_method,
            },
        )
        await self._db.commit()
        return PaymentRecordResponse.model_validate(payment)

    async def create_receipt_declaration(
        self, tenant_id: UUID, data: ContributionReceiptDeclarationCreate, *,
        declarant_user_id: UUID, declarant_role_code: str,
    ) -> ContributionReceiptDeclarationResponse:
        profile = await MembershipRepository(self._db).get_by_id(tenant_id, data.membership_profile_id)
        if profile is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Member profile not found")
        payload = data.model_dump()
        payload["received_at"] = payload["received_at"] or datetime.now(UTC)
        payload.update({
            "declarant_user_id": declarant_user_id,
            "declarant_role_code": declarant_role_code,
        })
        record = await self._repo.create_receipt_declaration(tenant_id, payload)
        await self._audit.record_event(
            tenant_id=tenant_id, actor_user_id=declarant_user_id,
            action="receipt_declaration_created", entity_type="contribution_receipt_declaration",
            entity_id=record.id, module_key="contributions",
            details={"membership_profile_id": record.membership_profile_id, "amount": str(record.amount)},
        )
        await self._db.commit()
        return ContributionReceiptDeclarationResponse.model_validate(record)

    async def update_receipt_declaration(
        self, tenant_id: UUID, declaration_id: UUID, data: ContributionReceiptDeclarationUpdate,
        *, declarant_user_id: UUID,
    ) -> ContributionReceiptDeclarationResponse:
        record = await self._repo.get_receipt_declaration(tenant_id, declaration_id)
        if record is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Receipt declaration not found")
        if record.declarant_user_id != declarant_user_id or record.status not in {
            ContributionReceiptStatus.draft.value, ContributionReceiptStatus.clarification_requested.value,
        }:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Receipt declaration cannot be edited")
        updated = await self._repo.update_receipt_declaration(
            tenant_id, declaration_id, data.model_dump(exclude_unset=True)
        )
        assert updated is not None
        await self._audit.record_event(
            tenant_id=tenant_id, actor_user_id=declarant_user_id,
            action="receipt_declaration_updated", entity_type="contribution_receipt_declaration",
            entity_id=updated.id, module_key="contributions", details={"changes": data.model_dump(exclude_unset=True)},
        )
        await self._db.commit()
        return ContributionReceiptDeclarationResponse.model_validate(updated)

    async def submit_receipt_declaration(
        self, tenant_id: UUID, declaration_id: UUID, *, declarant_user_id: UUID,
    ) -> ContributionReceiptDeclarationResponse:
        record = await self._repo.get_receipt_declaration(tenant_id, declaration_id)
        if record is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Receipt declaration not found")
        if record.declarant_user_id != declarant_user_id or record.status not in {
            ContributionReceiptStatus.draft.value, ContributionReceiptStatus.clarification_requested.value,
        }:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Receipt declaration cannot be submitted")
        updated = await self._repo.update_receipt_declaration(tenant_id, declaration_id, {
            "status": ContributionReceiptStatus.submitted.value, "submitted_at": datetime.now(UTC),
            "processing_note": None,
        })
        assert updated is not None
        await self._audit.record_event(
            tenant_id=tenant_id, actor_user_id=declarant_user_id,
            action="receipt_declaration_submitted", entity_type="contribution_receipt_declaration",
            entity_id=updated.id, module_key="contributions", details={},
        )
        await self._db.commit()
        return ContributionReceiptDeclarationResponse.model_validate(updated)

    async def process_receipt_declaration(
        self, tenant_id: UUID, declaration_id: UUID, data: ContributionReceiptDeclarationProcess,
        *, processor_user_id: UUID,
    ) -> ContributionReceiptDeclarationResponse:
        record = await self._repo.get_receipt_declaration(tenant_id, declaration_id)
        if record is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Receipt declaration not found")
        if record.status != ContributionReceiptStatus.submitted.value:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Only submitted declarations can be processed")
        if data.action in {"rejected", "clarification_requested", "cancelled"} and not data.note:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="A processing note is required")
        if data.action in {"validated", "partially_validated"}:
            if data.contribution_record_id is None:
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Contribution record is required")
            amount = data.processed_amount or record.amount
            if amount > record.amount:
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Processed amount cannot exceed declared amount")
            if data.action == "validated" and amount != record.amount:
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Validated amount must equal declared amount")
            contribution = await self._repo.get_contribution(tenant_id, data.contribution_record_id)
            if contribution is None or contribution.membership_profile_id != record.membership_profile_id:
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Contribution record does not belong to this member")
            payment = await self._repo.create_payment(tenant_id, {
                "contribution_record_id": contribution.id, "amount": amount, "currency": record.currency,
                "paid_at": datetime.now(UTC), "payment_method": record.payment_method,
                "reference": record.reference, "recorded_by": processor_user_id,
                "metadata_json": '{"source":"receipt_declaration"}',
            })
            await self._repo.update_contribution(tenant_id, contribution.id, {"paid_amount": contribution.paid_amount + amount})
            updated = await self._repo.update_receipt_declaration(tenant_id, declaration_id, {
                "status": data.action, "processed_at": datetime.now(UTC), "processed_by_user_id": processor_user_id,
                "processed_amount": amount, "processing_note": data.note, "contribution_record_id": contribution.id,
                "payment_record_id": payment.id,
            })
        else:
            updated = await self._repo.update_receipt_declaration(tenant_id, declaration_id, {
                "status": data.action, "processed_at": datetime.now(UTC), "processed_by_user_id": processor_user_id,
                "processing_note": data.note,
            })
        assert updated is not None
        await self._audit.record_event(
            tenant_id=tenant_id, actor_user_id=processor_user_id,
            action=f"receipt_declaration_{data.action}", entity_type="contribution_receipt_declaration",
            entity_id=updated.id, module_key="contributions",
            details={"processed_amount": str(updated.processed_amount) if updated.processed_amount else None},
        )
        await self._db.commit()
        return ContributionReceiptDeclarationResponse.model_validate(updated)

    async def list_receipt_declarations(
        self, tenant_id: UUID, *, declarant_user_id: UUID | None = None,
        membership_profile_id: UUID | None = None,
    ) -> list[ContributionReceiptDeclarationResponse]:
        rows = await self._repo.list_receipt_declarations(
            tenant_id, declarant_user_id=declarant_user_id, membership_profile_id=membership_profile_id
        )
        return [ContributionReceiptDeclarationResponse.model_validate(row) for row in rows]

    async def list_payments(
        self, tenant_id: UUID, contribution_id: UUID
    ) -> list[PaymentRecordResponse]:
        payments = await self._repo.list_payments_by_contribution(
            tenant_id, contribution_id
        )
        return [PaymentRecordResponse.model_validate(p) for p in payments]

    async def list_tenant_payments(self, tenant_id: UUID) -> list[PaymentRecordResponse]:
        payments = await self._repo.list_payments_by_tenant(tenant_id)
        return [PaymentRecordResponse.model_validate(p) for p in payments]

    async def get_summary(
        self, tenant_id: UUID, year: int | None = None
    ) -> dict:
        return await self._repo.get_tenant_contribution_summary(tenant_id, year)

    async def import_csv(
        self,
        tenant_id: UUID,
        content: bytes,
        *,
        dry_run: bool = False,
        actor_user_id: UUID | None = None,
    ) -> ImportResult:
        rows = parse_csv(content)
        member_repo = MembershipRepository(self._db)
        errors: list[ImportRowError] = []
        success_count = 0

        for i, row in enumerate(rows, start=2):
            row_errors: list[ImportRowError] = []

            member_code = row.get("member_code", "").strip()
            year_str = row.get("year", "").strip()
            expected_str = row.get("expected_amount", "0").strip()
            paid_str = row.get("paid_amount", "0").strip()
            currency = row.get("currency", "EUR").strip().upper()
            status_val = row.get("status", "pending").strip()

            if not member_code:
                row_errors.append(ImportRowError(row_number=i, column="member_code", message="member_code is required"))

            year: int | None = None
            try:
                year = int(year_str)
                if year < 2000 or year > 2100:
                    row_errors.append(ImportRowError(row_number=i, column="year", message="year must be between 2000 and 2100"))
            except (ValueError, TypeError):
                row_errors.append(ImportRowError(row_number=i, column="year", message="year must be an integer"))

            expected = Decimal("0")
            paid = Decimal("0")
            try:
                expected = Decimal(expected_str)
                if expected < 0:
                    row_errors.append(ImportRowError(row_number=i, column="expected_amount", message="expected_amount must be >= 0"))
            except Exception:
                row_errors.append(ImportRowError(row_number=i, column="expected_amount", message="expected_amount must be a valid number"))

            try:
                paid = Decimal(paid_str)
                if paid < 0:
                    row_errors.append(ImportRowError(row_number=i, column="paid_amount", message="paid_amount must be >= 0"))
            except Exception:
                row_errors.append(ImportRowError(row_number=i, column="paid_amount", message="paid_amount must be a valid number"))

            if status_val not in ("pending", "partial", "paid", "overdue", "waived"):
                row_errors.append(ImportRowError(row_number=i, column="status", message=f"Invalid status '{status_val}'"))

            if not row_errors:
                profile = await member_repo.get_by_member_code(tenant_id, member_code)
                if not profile:
                    row_errors.append(ImportRowError(
                        row_number=i, column="member_code", message=f"Unknown member_code '{member_code}'"
                    ))

            if row_errors:
                errors.extend(row_errors)
                continue

            if dry_run:
                success_count += 1
                continue

            assert profile is not None
            assert year is not None
            data = ContributionRecordCreate(
                membership_profile_id=profile.id,
                year=year,
                expected_amount=expected,
                paid_amount=paid,
                currency=currency,
                status=ContributionStatus(status_val),
            )
            record = await self._repo.create_contribution(tenant_id, data.model_dump())
            await self._audit.record_event(
                tenant_id=tenant_id,
                actor_user_id=actor_user_id,
                action="import",
                entity_type="contribution_record",
                entity_id=record.id,
                module_key="contributions",
                details={"member_code": member_code, "year": year, "source": "csv_import"},
            )
            success_count += 1

        if not dry_run:
            await self._db.commit()

        return ImportResult(
            total_rows=len(rows),
            success_count=success_count,
            error_count=len(errors),
            errors=errors,
            dry_run=dry_run,
        )

    async def export_csv(self, tenant_id: UUID) -> str:
        records = await self._repo.list_by_tenant(tenant_id)
        rows = [
            {
                "membership_profile_id": str(r.membership_profile_id),
                "year": str(r.year),
                "expected_amount": str(r.expected_amount),
                "paid_amount": str(r.paid_amount),
                "balance": str(r.balance),
                "currency": r.currency,
                "status": r.status,
                "due_date": str(r.due_date) if r.due_date else "",
            }
            for r in records
        ]
        return generate_csv(rows)

    async def export_finance_report_csv(self, tenant_id: UUID) -> str:
        records = await self._repo.list_by_tenant(tenant_id)
        payments = await self._repo.list_payments_by_tenant(tenant_id)

        payment_count_by_contribution: dict[str, int] = {}
        for payment in payments:
            key = str(payment.contribution_record_id)
            payment_count_by_contribution[key] = payment_count_by_contribution.get(key, 0) + 1

        rows = [
            {
                "contribution_id": str(record.id),
                "membership_profile_id": str(record.membership_profile_id),
                "year": str(record.year),
                "expected_amount": str(record.expected_amount),
                "paid_amount": str(record.paid_amount),
                "balance": str(record.balance),
                "currency": record.currency,
                "status": record.status,
                "payment_count": str(payment_count_by_contribution.get(str(record.id), 0)),
                "due_date": str(record.due_date) if record.due_date else "",
            }
            for record in records
        ]
        return generate_csv(rows)

    async def export_member_finance_report(
        self,
        tenant_id: UUID,
        *,
        year: int,
        export_format: str,
    ) -> tuple[bytes | str, str, str]:
        """Build a share-safe, human-readable finance export for audit roles."""
        tenant = await self._db.get(Tenant, tenant_id)
        profiles = await MembershipRepository(self._db).list_by_tenant(tenant_id)
        contributions = await self._repo.list_by_tenant(tenant_id, year)
        sanctions_result = await self._db.execute(
            select(DisciplinaryRecord).where(
                DisciplinaryRecord.tenant_id == tenant_id,
                DisciplinaryRecord.status.in_(("open", "under_review")),
            )
        )
        sanctions = list(sanctions_result.scalars().all())

        contribution_totals: dict[UUID, dict[str, Decimal]] = {}
        for record in contributions:
            totals = contribution_totals.setdefault(
                record.membership_profile_id,
                {"expected": Decimal("0"), "paid": Decimal("0"), "balance": Decimal("0")},
            )
            totals["expected"] += record.expected_amount
            totals["paid"] += record.paid_amount
            totals["balance"] += record.balance

        sanctions_by_member: dict[UUID, Decimal] = {}
        for sanction in sanctions:
            sanctions_by_member[sanction.membership_profile_id] = (
                sanctions_by_member.get(sanction.membership_profile_id, Decimal("0")) + sanction.amount
            )

        rows: list[dict[str, object]] = []
        for profile in sorted(profiles, key=lambda item: (item.last_name.lower(), item.first_name.lower())):
            totals = contribution_totals.get(
                profile.id,
                {"expected": Decimal("0"), "paid": Decimal("0"), "balance": Decimal("0")},
            )
            contribution_balance = totals["balance"]
            sanctions_due = sanctions_by_member.get(profile.id, Decimal("0"))
            if totals["expected"] == Decimal("0"):
                contribution_status = "Non enregistrée"
            elif contribution_balance <= Decimal("0"):
                contribution_status = "Soldée"
            elif totals["paid"] > Decimal("0"):
                contribution_status = "Partiellement payée"
            else:
                contribution_status = "À payer"
            rows.append(
                {
                    "member_code": profile.member_code,
                    "last_name": profile.last_name,
                    "first_name": profile.first_name,
                    "email": profile.email or "",
                    "phone": profile.phone or "",
                    "membership_type": "Famille" if profile.membership_type == "family" else "Individuel",
                    "membership_status": profile.status,
                    "expected": totals["expected"],
                    "paid": totals["paid"],
                    "contribution_balance": contribution_balance,
                    "contribution_status": contribution_status,
                    "sanctions_due": sanctions_due,
                    "total_due": contribution_balance + sanctions_due,
                }
            )

        organization_name = tenant.name if tenant else "Association"
        if export_format == "xlsx":
            return self._build_finance_xlsx(organization_name, year, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"rapport-financier-{year}.xlsx"
        if export_format == "pdf":
            return self._build_finance_pdf(organization_name, year, rows), "application/pdf", f"rapport-financier-{year}.pdf"
        if export_format == "whatsapp":
            return self._build_whatsapp_summary(organization_name, year, rows), "text/plain; charset=utf-8", f"rapport-financier-{year}-whatsapp.txt"
        raise ValueError(f"Unsupported finance export format: {export_format}")

    def _build_finance_xlsx(self, organization_name: str, year: int, rows: list[dict[str, object]]) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cotisations"
        sheet.sheet_view.showGridLines = False
        headers = ["Matricule", "Nom", "Prénom", "E-mail", "Téléphone", "Adhésion", "Statut membre", "Cotisation attendue (EUR)", "Cotisation versée (EUR)", "Reste cotisation (EUR)", "Statut cotisation", "Sanctions à payer (EUR)", "Total à payer (EUR)"]
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet["A1"] = f"{organization_name} - Rapport des cotisations {year}"
        sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="1F4F8F")
        sheet["A1"].alignment = Alignment(horizontal="center")
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        sheet["A2"] = "Document de suivi financier confidentiel - trésorerie et commissariat aux comptes"
        sheet["A2"].font = Font(italic=True, color="5B677A")
        sheet["A2"].alignment = Alignment(horizontal="center")
        sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append([
                row["member_code"],
                row["last_name"],
                row["first_name"],
                row["email"] or None,
                row["phone"] or None,
                row["membership_type"],
                row["membership_status"],
                float(row["expected"]),
                float(row["paid"]),
                float(row["contribution_balance"]),
                row["contribution_status"],
                float(row["sanctions_due"]),
                float(row["total_due"]),
            ])
        header_row = 4
        header_fill = PatternFill("solid", fgColor="DCE6F1")
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="17365D")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for column in (8, 9, 10, 12, 13):
            for cell in sheet.iter_cols(
                min_col=column,
                max_col=column,
                min_row=header_row + 1,
                max_row=header_row + len(rows),
            ):
                cell[0].number_format = '#,##0.00 "EUR"'
        widths = [18, 18, 18, 28, 18, 14, 16, 20, 20, 22, 20, 22, 18]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.freeze_panes = "A5"
        # Use a standard worksheet filter instead of an Excel Table.  Some mobile
        # spreadsheet readers repair the generated table XML and consequently make
        # the workbook appear corrupted, while a native filter is universally read.
        sheet.auto_filter.ref = f"A{header_row}:M{max(header_row, header_row + len(rows))}"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _build_finance_pdf(self, organization_name: str, year: int, rows: list[dict[str, object]]) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        output = io.BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
        styles = getSampleStyleSheet()
        story = [Paragraph(f"<b>{organization_name}</b>", styles["Title"]), Paragraph(f"Rapport des cotisations et sanctions - {year}", styles["Heading2"]), Paragraph("Document confidentiel destiné à la trésorerie et au commissariat aux comptes.", styles["BodyText"]), Spacer(1, 6 * mm)]
        header = ["Matricule", "Nom", "Prénom", "E-mail", "Téléphone", "Cotisation attendue", "Versée", "Reste", "Statut", "Sanctions", "Total dû"]
        data = [header]
        for row in rows:
            data.append([str(row["member_code"]), str(row["last_name"]), str(row["first_name"]), str(row["email"]), str(row["phone"]), f"{row['expected']:.2f} EUR", f"{row['paid']:.2f} EUR", f"{row['contribution_balance']:.2f} EUR", str(row["contribution_status"]), f"{row['sanctions_due']:.2f} EUR", f"{row['total_due']:.2f} EUR"])
        table = Table(data, repeatRows=1, colWidths=[30 * mm, 18 * mm, 18 * mm, 32 * mm, 21 * mm, 21 * mm, 19 * mm, 19 * mm, 24 * mm, 19 * mm, 19 * mm])
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4F8F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 6), ("LEADING", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2EC")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")]), ("ALIGN", (5, 1), (-1, -1), "RIGHT")]))
        story.append(table)
        document.build(story)
        return output.getvalue()

    def _build_whatsapp_summary(self, organization_name: str, year: int, rows: list[dict[str, object]]) -> str:
        lines = [
            f"*{organization_name}*",
            f"*Situation complète des cotisations - {year}*",
            f"{len(rows)} membre(s) inclus.",
            "",
        ]
        for index, row in enumerate(rows, start=1):
            lines.extend([f"*{index}. {row['first_name']} {row['last_name']}* ({row['member_code']})", f"Cotisation : {row['expected']:.2f} EUR | Versée : {row['paid']:.2f} EUR | Reste : {row['contribution_balance']:.2f} EUR", f"Sanctions à payer : {row['sanctions_due']:.2f} EUR | *Total dû : {row['total_due']:.2f} EUR*", f"Statut : {row['contribution_status']}", ""])
        return "\n".join(lines).strip() + "\n"

    async def list_reminders(
        self,
        tenant_id: UUID,
        *,
        year: int | None = None,
        profile_id: UUID | None = None,
        limit: int = 50,
    ) -> list[ContributionReminderResponse]:
        reminders = await self._repo.list_reminders_by_tenant(
            tenant_id,
            year=year,
            profile_id=profile_id,
            limit=limit,
        )
        return [ContributionReminderResponse.model_validate(reminder) for reminder in reminders]

    async def send_reminder(
        self,
        tenant_id: UUID,
        contribution_id: UUID,
        payload: ContributionReminderSendRequest,
        *,
        actor_user_id: UUID | None = None,
    ) -> ContributionReminderResponse:
        contribution = await self._repo.get_contribution(tenant_id, contribution_id)
        if not contribution:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Contribution record not found",
            )

        member_repo = MembershipRepository(self._db)
        profile = await member_repo.get_by_id(tenant_id, contribution.membership_profile_id)
        if not profile:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Member profile not found for this contribution",
            )

        reminder = await self._dispatch_reminder(
            tenant_id,
            contribution=contribution,
            profile=profile,
            channel=payload.channel,
            actor_user_id=actor_user_id,
        )
        await self._db.commit()
        return ContributionReminderResponse.model_validate(reminder)

    async def send_batch_reminders(
        self,
        tenant_id: UUID,
        payload: ContributionReminderBatchRequest,
        *,
        actor_user_id: UUID | None = None,
    ) -> ContributionReminderBatchResponse:
        contributions = await self._repo.list_by_tenant(tenant_id, payload.year)
        selected = self._filter_reminder_candidates(contributions, payload)
        if len(selected) > payload.limit:
            selected = selected[: payload.limit]

        member_repo = MembershipRepository(self._db)
        reminders = []
        for contribution in selected:
            profile = await member_repo.get_by_id(tenant_id, contribution.membership_profile_id)
            if not profile:
                continue
            reminder = await self._dispatch_reminder(
                tenant_id,
                contribution=contribution,
                profile=profile,
                channel=payload.channel,
                actor_user_id=actor_user_id,
            )
            reminders.append(ContributionReminderResponse.model_validate(reminder))

        await self._db.commit()
        return ContributionReminderBatchResponse(
            attempted_count=len(selected),
            reminder_count=len(reminders),
            reminders=reminders,
        )

    def _filter_reminder_candidates(
        self,
        contributions: list,
        payload: ContributionReminderBatchRequest,
    ) -> list:
        now = datetime.now(UTC)
        due_soon_cutoff = now + timedelta(days=14)
        filtered = []
        for contribution in contributions:
            if Decimal(str(contribution.balance)) <= Decimal("0.00"):
                continue
            if contribution.status in {"paid", "waived"}:
                continue
            if payload.status is not None and contribution.status != payload.status.value:
                continue
            due_date = contribution.due_date
            if payload.due_scope == "overdue":
                if not (
                    contribution.status == "overdue"
                    or (due_date is not None and due_date < now)
                ):
                    continue
            elif payload.due_scope == "due_soon":
                if due_date is None or due_date > due_soon_cutoff:
                    continue
            filtered.append(contribution)
        return filtered

    async def _dispatch_reminder(
        self,
        tenant_id: UUID,
        *,
        contribution,
        profile: MembershipProfile,
        channel: str,
        actor_user_id: UUID | None,
    ):
        subject, body = self._build_reminder_message(profile, contribution)
        recipient = (profile.email or "").strip()

        if Decimal(str(contribution.balance)) <= Decimal("0.00") or contribution.status in {"paid", "waived"}:
            return await self._create_reminder_record(
                tenant_id,
                contribution=contribution,
                profile=profile,
                channel=channel,
                subject=subject,
                body=body,
                recipient=recipient or "missing-recipient",
                delivery_status=ReminderDeliveryStatus.skipped.value,
                provider_message="Reminder skipped because the contribution has no outstanding balance.",
                actor_user_id=actor_user_id,
                audit_action="reminder_skipped",
            )

        if channel != "email":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Only email reminders are supported in this sprint",
            )

        if not recipient:
            return await self._create_reminder_record(
                tenant_id,
                contribution=contribution,
                profile=profile,
                channel=channel,
                subject=subject,
                body=body,
                recipient="missing-email",
                delivery_status=ReminderDeliveryStatus.skipped.value,
                provider_message="Reminder skipped because the member has no email address.",
                actor_user_id=actor_user_id,
                audit_action="reminder_skipped",
            )

        provider = next(
            (item for item in self._notification_providers if getattr(item, "channel", "") == channel),
            None,
        )
        if provider is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"No notification provider is available for channel '{channel}'",
            )

        try:
            dispatched = await provider.send_message(
                tenant_id=tenant_id,
                actor_user_id=actor_user_id,
                recipient=recipient,
                subject=subject,
                body=body,
            )
        except Exception as exc:  # pragma: no cover - defensive safeguard
            dispatched = NotificationDispatchResult(
                channel=channel,
                status=ReminderDeliveryStatus.failed.value,
                message=f"Reminder delivery failed: {exc}",
                delivered=False,
                simulation_only=False,
            )

        audit_action = (
            "reminder_sent"
            if dispatched.status in {ReminderDeliveryStatus.sent.value, ReminderDeliveryStatus.simulated.value}
            else "reminder_failed"
        )
        return await self._create_reminder_record(
            tenant_id,
            contribution=contribution,
            profile=profile,
            channel=channel,
            subject=subject,
            body=body,
            recipient=recipient,
            delivery_status=dispatched.status,
            provider_message=dispatched.message,
            actor_user_id=actor_user_id,
            audit_action=audit_action,
        )

    async def _create_reminder_record(
        self,
        tenant_id: UUID,
        *,
        contribution,
        profile: MembershipProfile,
        channel: str,
        subject: str,
        body: str,
        recipient: str,
        delivery_status: str,
        provider_message: str | None,
        actor_user_id: UUID | None,
        audit_action: str,
    ):
        reminder = await self._repo.create_reminder(
            tenant_id,
            {
                "contribution_record_id": contribution.id,
                "membership_profile_id": profile.id,
                "member_display_name": profile.display_name,
                "member_code": profile.member_code,
                "balance_snapshot": contribution.balance,
                "due_date_snapshot": contribution.due_date,
                "channel": channel,
                "delivery_status": delivery_status,
                "recipient": recipient,
                "subject": subject,
                "body": body,
                "provider_message": provider_message,
                "reminded_by": actor_user_id,
            },
        )
        await self._audit.record_event(
            tenant_id=tenant_id,
            actor_user_id=actor_user_id,
            action=audit_action,
            entity_type="contribution_reminder",
            entity_id=reminder.id,
            module_key="contributions",
            details={
                "contribution_record_id": contribution.id,
                "membership_profile_id": profile.id,
                "channel": channel,
                "delivery_status": delivery_status,
                "recipient": recipient,
            },
        )
        return reminder

    def _build_reminder_message(self, profile: MembershipProfile, contribution) -> tuple[str, str]:
        due_line = ""
        if contribution.due_date is not None:
            due_line = f" The recorded due date is {contribution.due_date.strftime('%Y-%m-%d')}."
        subject = f"Contribution reminder for {contribution.year}"
        body = (
            f"Hello {profile.display_name},\n\n"
            f"This is a reminder that your contribution balance for {contribution.year} is "
            f"{Decimal(str(contribution.balance)).quantize(Decimal('0.01'))} {contribution.currency}."
            f"{due_line}\n\n"
            "If you have already paid or need clarification, please contact the treasury team of your organization.\n\n"
            "Kairo automated reminder"
        )
        return subject, body
